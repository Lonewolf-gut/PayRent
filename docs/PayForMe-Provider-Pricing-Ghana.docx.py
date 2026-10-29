#!/usr/bin/env python3
"""Generate PayForMe provider pricing Word document."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "/workspace/docs/PayForMe-Provider-Pricing-Ghana.docx"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "1F4E79") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], header_fill)
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, text in enumerate(row):
            cells[c_idx].text = text
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    title = doc.add_heading("PayForMe — Ghana Provider Pricing Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("SMS · USSD · Online payments · Disbursements")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True
    doc.add_paragraph("Date: October 2026")
    doc.add_paragraph(
        "Costs scale with usage. Figures marked tax exclusive exclude VAT/NHIL unless stated. "
        "Confirm final rates with each vendor before signing."
    )

    doc.add_heading("1. SMS providers", level=1)
    doc.add_heading("1.1 Arkesel (PayForMe default)", level=2)
    add_table(
        doc,
        ["Service", "Cost"],
        [
            ["Account / integration", "GHS 0/month"],
            ["Phone OTP (SMS + USSD)", "GHS 0.035 per verification"],
            ["Bulk SMS", "GHS 0.022–0.031 per SMS"],
        ],
    )
    doc.add_paragraph("Example: 500 verifications → 500 × GHS 0.035 = GHS 17.50")

    doc.add_heading("1.2 Hubtel SMS", level=2)
    add_table(
        doc,
        ["Service", "Cost"],
        [
            ["Account / integration", "GHS 0/month"],
            ["SMS / OTP", "GHS 0.007–0.025 per SMS"],
        ],
    )

    doc.add_heading("1.3 Wigal FROG SMS (tax exclusive)", level=2)
    add_table(
        doc,
        ["Monthly SMS volume", "Price per SMS"],
        [
            ["0 – 100,000", "GHS 0.0300"],
            ["100,001 – 500,000", "GHS 0.0280"],
            ["500,001 – 1,000,000", "GHS 0.0250"],
            ["Above 1,000,000", "GHS 0.0220"],
        ],
    )
    doc.add_paragraph("Example: 500 SMS → 500 × GHS 0.03 = GHS 15.00 (tax exclusive)")

    doc.add_page_break()
    doc.add_heading("2. USSD — Wigal (tax exclusive)", level=1)
    doc.add_paragraph(
        "Monthly traffic charge applies across all networks. Each menu step = one hit (GHS 0.02). "
        "Fund FROG wallet for usage; rental invoiced monthly."
    )
    add_table(
        doc,
        ["Plan", "Shared code", "Example", "Rental / month", "Usage / hit"],
        [
            ["Starter", "5 digits", "*800*12345#", "GHS 70", "GHS 0.02"],
            ["Basic", "4 digits", "*800*1234#", "GHS 100", "GHS 0.02"],
            ["Standard (recommended)", "3 digits", "*800*123#", "GHS 150", "GHS 0.02"],
            ["Premium", "2 digits", "*800*12#", "GHS 300", "GHS 0.02"],
            ["Business", "1 digit", "*800*1#", "GHS 800", "GHS 0.02"],
        ],
    )
    doc.add_heading("USSD development (one-time)", level=2)
    add_table(
        doc,
        ["Item", "Cost"],
        [
            ["Basic USSD short code app", "GHS 1,000 – 2,000"],
            ["USSD + API / banking integration", "Custom quote — request fixed price"],
        ],
    )
    doc.add_paragraph("Example USSD monthly (Standard plan, 2,000 sessions × 5 hits):")
    doc.add_paragraph("Rental GHS 150 + hits 2,000 × 5 × GHS 0.02 = GHS 200 → Total GHS 350 (tax exclusive)")

    doc.add_page_break()
    doc.add_heading("3. Paystack Ghana (integrated)", level=1)
    add_table(
        doc,
        ["Service", "Cost"],
        [
            ["Account / integration", "GHS 0/month"],
            ["Local payments", "1.95% per transaction"],
            ["International payments", "1.95%"],
            ["Transfer to Mobile Money", "GHS 1 per successful transfer"],
            ["Transfer to bank account", "GHS 8 per successful transfer"],
        ],
    )
    doc.add_paragraph("Example: GHS 100,000 collections → GHS 1,950 fees (not a monthly subscription).")
    doc.add_paragraph("500 MoMo payouts → 500 × GHS 1 = GHS 500. Total with collections: GHS 2,450.")

    doc.add_heading("4. Hubtel Ghana (integrated)", level=1)
    add_table(
        doc,
        ["Service", "Cost"],
        [
            ["Account / integration", "GHS 0/month"],
            ["MoMo collections", "1.95% (min ~GHS 0.30)"],
            ["Send Money payouts", "Confirm with Hubtel sales (not public)"],
        ],
    )
    doc.add_paragraph("Collections example: GHS 100,000 × 1.95% = GHS 1,950.")

    doc.add_heading("5. MTN MoMo API (integrated)", level=1)
    add_table(
        doc,
        ["Service", "Cost"],
        [
            ["Collections (published default)", "2% per transaction (negotiable at volume)"],
            ["API calls", "GHS 0 per call"],
        ],
    )

    doc.add_page_break()
    doc.add_heading("6. Cheapest online provider PayForMe already has", level=1)
    add_table(
        doc,
        ["Provider", "Integrated?", "Collection fee", "Payouts (published)"],
        [
            ["Paystack", "Yes", "1.95%", "GHS 1 MoMo / GHS 8 bank"],
            ["Hubtel", "Yes", "1.95%", "Quote required"],
            ["MTN MoMo direct", "Yes", "2% default", "Telco fees apply"],
        ],
    )
    doc.add_paragraph(
        "Recommendation: Paystack for default online payments — same 1.95% as Hubtel on collections, "
        "with published low payout fees. MTN direct only if negotiated below 1.95% at high volume."
    )

    doc.add_heading("7. Paystack vs Hubtel", level=1)
    add_table(
        doc,
        ["Item", "Paystack", "Hubtel"],
        [
            ["Monthly subscription", "GHS 0", "GHS 0"],
            ["GHS 100,000 collections", "GHS 1,950", "GHS 1,950"],
            ["500 MoMo payouts", "GHS 500", "Sales quote"],
            ["500 bank payouts", "GHS 4,000", "Sales quote"],
        ],
    )

    doc.add_heading("8. Suggested stacks", level=1)
    add_table(
        doc,
        ["Stack", "Best for"],
        [
            ["Arkesel SMS + Paystack payments", "Current PayForMe — clear payout pricing"],
            ["Wigal SMS + Wigal USSD Standard + Redde/Paystack", "Client USSD + offline menus"],
            ["Hubtel SMS + Hubtel payments", "Single vendor"],
        ],
    )

    doc.add_heading("9. Contacts", level=1)
    doc.add_paragraph("Arkesel: https://arkesel.com/pricing/")
    doc.add_paragraph("Paystack Ghana: https://website-v4.paystack.co/gh/pricing")
    doc.add_paragraph("Hubtel: https://developers.hubtel.com/")
    doc.add_paragraph("Wigal: support@wigal.com.gh · +233 55 301 9529 · frog.wigal.com.gh")

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
